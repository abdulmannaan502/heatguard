"""Fill the official submission template with HeatGuard content (template layout and headings are kept)."""
import sys, os, copy, json
ROOT = os.path.join(os.path.dirname(__file__), "..")
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree

TPL = sys.argv[1]; OUT = sys.argv[2]
EV = json.load(open(f"{ROOT}/results/evaluation.json"))["holdout_202"]
REP = json.load(open(f"{ROOT}/results/evaluation.json"))["replicate_303"]
AB = json.load(open(f"{ROOT}/results/ablation.json"))
INK, BROWN, TAUPE, CREAM = RGBColor(0x3b, 0x31, 0x2b), RGBColor(0x7a, 0x66, 0x56), RGBColor(0xd9, 0xcf, 0xc1), RGBColor(0xf0, 0xe9, 0xdf)
GREEN, AMB, RED = RGBColor(0x2f, 0x7a, 0x57), RGBColor(0xc9, 0x8a, 0x12), RGBColor(0xbf, 0x3b, 0x2b)
FONT = "Outfit"
prs = Presentation(TPL)
S = list(prs.slides)

def para_fmt(p, bullet=False, size=10, space=3):
    pPr = p._p.get_or_add_pPr()
    for ch in list(pPr): pPr.remove(ch)
    if bullet:
        pPr.set("marL", "150000"); pPr.set("indent", "-150000")
        sb = etree.SubElement(pPr, qn("a:spcBef")); etree.SubElement(sb, qn("a:spcPts")).set("val", str(space * 100))
        bu = etree.SubElement(pPr, qn("a:buChar")); bu.set("char", "•")
    else:
        pPr.set("marL", "0"); pPr.set("indent", "0")
        sb = etree.SubElement(pPr, qn("a:spcBef")); etree.SubElement(sb, qn("a:spcPts")).set("val", str(space * 100))
        etree.SubElement(pPr, qn("a:buNone"))

def add_runs(p, parts, size, color=INK):
    """parts: str or list of (text, bold) / (text, bold, color)"""
    if isinstance(parts, str): parts = [(parts, False)]
    for part in parts:
        r = p.add_run(); r.text = part[0]; r.font.size = Pt(size); r.font.bold = part[1]; r.font.name = FONT
        r.font.color.rgb = part[2] if len(part) > 2 else color

def tb(slide, x, y, w, h, items, size=10, color=INK, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, margin=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = box.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin); tf.margin_top = tf.margin_bottom = Inches(0.02)
    first = True
    for it in items:
        if isinstance(it, dict): parts, bullet, sz, sp = it["t"], it.get("b", False), it.get("size", size), it.get("sp", 3)
        else: parts, bullet, sz, sp = it, False, size, 3
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        para_fmt(p, bullet, sz, sp if not first else 0); p.alignment = align
        add_runs(p, parts, sz, color)
    return box

def B(parts, size=None, sp=3):
    d = {"t": parts, "b": True, "sp": sp}
    if size: d["size"] = size
    return d

def H(text, size=10.5, sp=6):
    return {"t": [(text, True, BROWN)], "size": size, "sp": sp}

def rect(slide, x, y, w, h, text, fill=CREAM, line=BROWN, size=9, bold=False, color=INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, align=PP_ALIGN.CENTER):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(0.9)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = 0.12
    sh.shadow.inherit = False
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.02)
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; para_fmt(p, False, size, 0)
        if isinstance(ln, tuple): add_runs(p, [(ln[0], ln[1])] if len(ln) == 2 else [ln], ln[2] if len(ln) > 2 and False else size, color)
        else: add_runs(p, [(ln, bold)], size, color)
    return sh

def arrow(slide, x1, y1, x2, y2, color=BROWN):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(1.4)
    ln = c.line._get_or_add_ln(); te = etree.SubElement(ln, qn("a:tailEnd")); te.set("type", "triangle")
    return c

def table(slide, x, y, w, rows, colw, size=8.5, header=True, rowh=0.24):
    gt = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(rowh * len(rows)))
    t = gt.table
    tblPr = t._tbl.tblPr
    for a in ("firstRow", "bandRow"): tblPr.set(a, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None: sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # No Style, No Grid
    for j, cw in enumerate(colw): t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(rowh)
        for j, val in enumerate(row):
            c = t.cell(i, j); c.margin_left = c.margin_right = Inches(0.05); c.margin_top = c.margin_bottom = Inches(0.025)
            c.fill.solid(); c.fill.fore_color.rgb = BROWN if (header and i == 0) else (CREAM if i % 2 else RGBColor(0xfb, 0xf8, 0xf3))
            tf = c.text_frame; tf.word_wrap = True; p = tf.paragraphs[0]; para_fmt(p, False, size, 0)
            add_runs(p, [(str(val), header and i == 0 or j == 0, RGBColor(255, 255, 255) if (header and i == 0) else INK)], size)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
    return gt

# ---------------------------------------------------------------- slide 1
s = S[0]
box = [sh for sh in s.shapes if sh.has_text_frame][0]
paras = box.text_frame.paragraphs
def set_val(p, label, value, size=14):
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    r0 = p.runs[0]; r0.text = label; r0.font.size = Pt(size)
    r1 = copy.deepcopy(r0._r); p._p.insert(list(p._p).index(r0._r) + 1, r1)
    from pptx.text.text import _Run
    run = _Run(r1, p); run.text = value; run.font.bold = True; run.font.size = Pt(size)
set_val(paras[0], "Team Name: ", "HeatGuard (solo entry, Abdul Mannaan Sayyad, Triad AI)")
set_val(paras[1], "Team Size: ", "1")
set_val(paras[2], "Problem Statement: ", "[Population Health] Climate-Resilient Worker Health. Outdoor workers collapse from heat because weather alerts and fixed midday breaks treat everyone the same, while real risk is individual. HeatGuard estimates each worker's core temperature from a wearable and the weather, forecasts it 30 minutes ahead, and sends a tiered alert in the worker's own language.", 11)
paras[2].runs[0].font.size = Pt(11)
for r in paras[2].runs[1:]: r.font.bold = False

# ---------------------------------------------------------------- slide 2
s = S[1]
gt = [sh for sh in s.shapes if sh.has_table][0]; t = gt.table
vals = ["1", "Abdul Mannaan Sayyad", "India", "AI Engineer (team lead, sole member)", "Triad AI", "Confirmed: age 21, over 18"]
for j, v in enumerate(vals):
    c = t.cell(1, j); tf = c.text_frame; p = tf.paragraphs[0]
    for r in list(p.runs): r._r.getparent().remove(r._r)
    r = p.add_run(); r.text = v; r.font.size = Pt(9); r.font.name = FONT
for i in range(2, 6):
    for j in range(1, 6): pass
tb(s, 0.4, 4.65, 9.2, 0.8, [
    {"t": [("Solo entry. ", True, BROWN), ("One person designed, built, validated and documented the whole solution. No other member is listed, so rows 2 to 5 are intentionally left empty.", False)], "size": 9},
    {"t": [("Eligibility. ", True, BROWN), ("The team lead is 18 or older. Permission from the employer or institution (Triad AI) is the participant's responsibility under the challenge rules and is being confirmed by the applicant.", False)], "size": 9}])

# ---------------------------------------------------------------- slide 3 capability
s = S[2]
cap = ("Solo AI engineer (21) at Triad AI, India. For this challenge I built the full HeatGuard stack myself: a physics-based worker and weather simulator, an extended Kalman filter that estimates core temperature from wearable heart rate, a 30-minute forecaster, work-rest policies, a closed-loop evaluation on 6,000 synthetic worker-shifts, a multilingual alert layer, a privacy-by-design API with tests, and a live site console. "
       "Skills used: machine learning and state estimation, signal processing, Python and FastAPI, data visualisation, responsible-AI design. "
       "Gaps I am addressing: no clinician or occupational-health physician on the team yet, no human data yet, and translations are machine-drafted. Next step is to recruit an occupational-health advisor and a construction partner for a calibration study and pilot.")
print("capability chars", len(cap))
tb(s, 0.35, 1.5, 9.3, 2.4, [{"t": cap, "size": 11.5, "sp": 0}])
tb(s, 0.35, 4.15, 9.3, 0.9, [{"t": [("Built and tested in this challenge: ", True, BROWN), ("Python package with 7 automated tests, FastAPI service, replay console, 3-minute demo video. All results in this deck are reproducible from the code.", False)], "size": 10}])

# ---------------------------------------------------------------- slide 4 architecture
s = S[3]
y0 = 1.55
rect(s, 0.3, y0, 2.05, 0.5, ["Wearable heart rate", "any wrist band or chest strap"], size=8.5)
rect(s, 0.3, y0 + 0.62, 2.05, 0.5, ["Site weather", "air temp, humidity, sun, wind"], size=8.5)
rect(s, 0.3, y0 + 1.24, 2.05, 0.5, ["Task class", "light to very heavy"], size=8.5)
rect(s, 0.3, y0 + 1.86, 2.05, 0.5, ["Days on the job", "heat acclimatization"], size=8.5)
tb(s, 0.3, y0 - 0.2, 2.05, 0.2, [{"t": [("INPUTS (consented)", True, BROWN)], "size": 8}])
core = rect(s, 2.85, y0, 2.35, 2.36, [("HeatGuard engine", True), "", "Extended Kalman filter", "core + skin temperature", "", "Heat-balance model:", "metabolism, sweat, evaporation,", "hydration, acclimatization", "", "Heart-rate correction", "with outlier gating"], fill=BROWN, size=8.5, color=RGBColor(255, 255, 255))
for k in range(4): arrow(s, 2.35, y0 + 0.25 + 0.62 * k, 2.85, y0 + 0.4 + 0.5 * k)
rect(s, 5.7, y0 + 0.15, 1.75, 0.95, ["30-minute forecast", "with uncertainty", "Tc now and Tc in 30 min"], size=8.5)
rect(s, 5.7, y0 + 1.25, 1.75, 0.95, ["Tiering and work-rest rule", "GREEN / AMBER / RED", "human first-aider on RED"], size=8.5)
arrow(s, 5.2, y0 + 0.62, 5.7, y0 + 0.62); arrow(s, 5.2, y0 + 1.72, 5.7, y0 + 1.72); arrow(s, 6.58, y0 + 1.1, 6.58, y0 + 1.25)
tb(s, 7.85, y0 - 0.2, 1.9, 0.2, [{"t": [("OUTPUTS", True, BROWN)], "size": 8}])
rect(s, 7.85, y0, 1.85, 0.7, ["Worker phone", "own language, WhatsApp / SMS", "owns the data"], size=8)
rect(s, 7.85, y0 + 0.82, 1.85, 0.7, ["Supervisor", "risk tier and action only", "no temperature or heart rate"], size=8)
rect(s, 7.85, y0 + 1.64, 1.85, 0.72, ["Health authority", "aggregated, k ≥ 5", "heat early-warning map"], size=8)
for k in range(3): arrow(s, 7.45, y0 + 1.7, 7.85, y0 + 0.35 + 0.82 * k)
rect(s, 0.3, y0 + 2.55, 9.4, 0.5, [("Privacy by design: ", True), "explicit consent, data minimisation, worker-owned data, role-based views, aggregate-only reporting, no use for hiring, pay or dismissal"], fill=RGBColor(0xdc, 0xeb, 0xe2), line=GREEN, size=8.5, color=INK)
tb(s, 0.3, y0 + 3.12, 9.4, 0.5, [{"t": [("Validation loop (built): ", True, BROWN), ("a hidden, individually varied synthetic worker (richer physiology, clothing, wind, sensor errors) feeds the same estimator that would run in the field, so the estimator never sees the truth. Screening and decision support only; not a diagnosis.", False)], "size": 8.5}])

# ---------------------------------------------------------------- slide 5 users / outcome
s = S[4]
tb(s, 0.35, 1.5, 4.5, 3.9, [
    H("Who it is for"),
    B([("Outdoor workers ", True), ("in construction, logistics, delivery and agriculture in hot climates, above all people in their first week on the job, when acclimatization is lowest.", False)]),
    B([("Site supervisors and safety officers ", True), ("who need a clear tier and action, not medical data.", False)]),
    B([("Health authorities and insurers ", True), ("who need an aggregated early-warning picture of heat illness risk.", False)]),
    H("Why current practice falls short"),
    B("Statutory midday breaks (UAE: 12:30 to 15:00, 15 June to 15 September) are a fixed time window. Strain also peaks in humid mornings and late afternoons."),
    B("Site-wide WBGT cut-offs treat a day-one rebar fixer and a ten-year veteran the same."),
    B("ILO estimates heat stress will cost 2.2% of global working hours by 2030, equal to 80 million full-time jobs."),
], size=9.5)
tb(s, 5.1, 1.5, 4.6, 3.9, [
    H("Intended outcome"),
    B([("Act before collapse: ", True), ("each worker gets a 30-minute-ahead warning and a specific action in their own language.", False)]),
    B([("Protect without stopping work: ", True), ("in simulation, ", False), (f"{EV['C']['exc385']:.1f}% of workers reach 38.5 °C with HeatGuard versus {EV['A']['exc385']:.1f}% today and {EV['B']['exc385']:.1f}% with a site-wide cut-off,", True), (f" while keeping {EV['C']['work_h']:.1f} productive hours per shift ({EV['B']['work_h']:.1f} h with the cut-off).", False)]),
    B([("Fair by design: ", True), ("the new-arrival group is protected most (50.8% to 3.9% in simulation).", False)]),
    B([("Worker-first data: ", True), ("the worker owns the data; the employer sees the tier only.", False)]),
    H("Setting"),
    B("First deployment: Abu Dhabi construction and delivery workforce, then any hot-climate outdoor workforce."),
], size=9.5)

# ---------------------------------------------------------------- slide 6 success
s = S[5]
tb(s, 0.35, 1.5, 4.7, 3.9, [
    H("Success measures for the pilot"),
    B([("Primary: ", True), ("share of worker-shifts reaching a core temperature of 38.5 °C (validated sensor or ingestible pill in a sub-study) and heat-illness incidents, against a matched baseline period.", False)]),
    B([("Estimate quality: ", True), ("RMSE ≤ 0.5 °C against pill reference; 95% interval coverage.", False)]),
    B([("Timeliness: ", True), ("at least 80% of 38.5 °C events warned 10 or more minutes ahead.", False)]),
    B([("Burden: ", True), ("RED alerts per worker-shift, alert acknowledgement rate, false-alarm rate.", False)]),
    B([("Productivity: ", True), ("productive hours kept versus status quo and cut-off rules.", False)]),
    B([("Trust and fairness: ", True), ("worker trust survey, and every metric split by acclimatization, age, BMI and language.", False)]),
    B([("Adoption: ", True), ("share of RED alerts followed by a rest within 5 minutes.", False)]),
], size=9.3)
rows = [["Held-out synthetic cohort (n = 3,000)", "A Status quo", "B WBGT cut-off", "C HeatGuard"],
        ["Reach 38.5 °C", f"{EV['A']['exc385']:.1f}%", f"{EV['B']['exc385']:.1f}%", f"{EV['C']['exc385']:.1f}%"],
        ["Reach 39.0 °C", f"{EV['A']['exc390']:.1f}%", f"{EV['B']['exc390']:.1f}%", f"{EV['C']['exc390']:.1f}%"],
        ["Productive hours / shift", f"{EV['A']['work_h']:.1f}", f"{EV['B']['work_h']:.1f}", f"{EV['C']['work_h']:.1f}"],
        ["Replicate seed, reach 38.5 °C", f"{REP['A']['exc385']:.1f}%", f"{REP['B']['exc385']:.1f}%", f"{REP['C']['exc385']:.1f}%"]]
table(s, 5.2, 1.55, 4.5, rows, [1.75, 0.9, 0.95, 0.9], size=8.5, rowh=0.28)
s.shapes.add_picture(f"{ROOT}/results/fig_outcomes.png", Inches(5.2), Inches(3.05), width=Inches(4.5))
tb(s, 5.2, 4.95, 4.5, 0.5, [{"t": "Synthetic validation. 95% CI for HeatGuard: " + f"{EV['C']['exc385_ci'][0]:.1f} to {EV['C']['exc385_ci'][1]:.1f}%. Not human data; real-world effect must be shown in the pilot.", "size": 8}])

# ---------------------------------------------------------------- slide 7 evidence
s = S[6]
tb(s, 0.35, 1.5, 4.75, 3.9, [
    H("Published evidence behind the approach"),
    B([("Heart rate can estimate core temperature with a Kalman filter. ", True), ("Buller et al. 2013, Physiological Measurement 34(6):781 (US Army); Falcone, Del Ferraro, Molinaro, Zollo, Lenzuni 2024, Frontiers in Public Health, doi 10.3389/fpubh.2024.1219595 (occupational use).", False)]),
    B([("Heat balance with core and skin nodes. ", True), ("Gagge et al. 1971, ASHRAE Transactions 77 (two-node model); respiratory heat-loss terms as in ISO 7933 style heat-strain models.", False)]),
    B([("Occupational limits. ", True), ("Core temperature of about 38.0 to 38.5 °C is commonly used as an upper guideline (NIOSH Criteria for a Recommended Standard, 2016; ISO 9886).", False)]),
    B([("Scale of the problem. ", True), ("ILO, Working on a warmer planet, 2019: 2.2% of working hours and 80 million full-time jobs at risk by 2030.", False)]),
    B([("Existing rule. ", True), ("UAE midday break, 12:30 to 15:00, 15 June to 15 September, fines up to Dh50,000 (Gulf News, June 2025).", False)]),
], size=8.8)
tb(s, 5.3, 1.5, 4.4, 3.9, [
    H("Our own testing (synthetic, reproducible)"),
    B([("Estimator: ", True), (f"RMSE {EV['estimator']['rmse']:.2f} °C over 3,000 workers; heart-rate correction cuts RMSE from {AB['no heart-rate correction (model only)']['rmse']:.2f} (model only) to {AB['full']['rmse']:.2f} °C.", False)]),
    B([("Early warning: ", True), (f"{EV['warning']['warned_10']:.0f}% of 38.5 °C events warned 10+ min ahead; median lead {EV['warning']['median_lead_min']:.0f} min.", False)]),
    B([("Robustness: ", True), (f"heart-rate sensitivity prior wrong by ±30% moves RMSE only to {AB['HR sensitivity prior -30 %']['rmse']:.2f} / {AB['HR sensitivity prior +30 %']['rmse']:.2f} °C; 4% sensor dropouts and 1.5% motion outliers included.", False)]),
    B([("Fairness: ", True), ("RMSE 0.40 to 0.44 °C across age and BMI bands; new arrivals 3.9% vs 4.5% for acclimatized workers reach 38.5 °C.", False)]),
    H("Limitation we state openly", sp=6),
    B("All numbers above come from a simulator built from published physiology, not from people. The simulator and estimator share a physics family, so absolute accuracy in the field will be lower. Calibration against ingestible-pill data is the first pilot task."),
], size=8.8)

# ---------------------------------------------------------------- slide 8 data
s = S[7]
rows = [["Source", "Type", "Owner", "Sensitivity", "Permission / legal basis", "Personal data?"],
        ["Synthetic worker-shifts (6,000 + demo)", "Simulated physiology, weather, sensors", "Applicant (generated by own code)", "None", "Own work, no third-party rights", "No"],
        ["Published equations and parameters", "Literature (Gagge, Stull, Tetens, DuBois, ISO-style terms)", "Public domain / cited authors", "None", "Citation of published science, no data copied", "No"],
        ["Weather (pilot)", "Site sensor or public weather service", "Site operator / provider", "Low", "Provider terms; site agreement", "No"],
        ["Wearable heart rate (pilot)", "Time series, 1 per minute", "The worker", "High (health data)", "Explicit informed consent, revocable; UAE PDPL (Federal Decree-Law 45/2021) and Abu Dhabi health-data rules to be confirmed with the regulator", "Yes, pseudonymised ID"],
        ["Reference core temperature (calibration study)", "Ingestible pill or validated sensor", "Study participants", "High", "Ethics approval and consent before any collection", "Yes, study only"]]
table(s, 0.3, 1.5, 9.4, rows, [1.7, 1.7, 1.4, 1.0, 2.5, 1.1], size=7.8, rowh=0.52)
tb(s, 0.3, 4.85, 9.4, 0.55, [{"t": [("No personal data was used in this submission. ", True, BROWN), ("Everything shown is synthetic. In a pilot, raw heart rate stays on the worker's phone or in a worker-keyed store; employers receive only the tier; authorities receive aggregates with a minimum group size of 5.", False)], "size": 8.8}])

# ---------------------------------------------------------------- slide 9 risks
s = S[8]
rows = [["Risk", "What could go wrong", "Safeguard in the design"],
        ["Privacy", "Heart rate and heat history reveal health", "Consent, minimisation, worker-owned data, supervisor sees tier only (tested), k ≥ 5 aggregates"],
        ["Security", "Leak of health-linked data", "Pseudonymous IDs, encryption in transit and at rest, role-based access, audit log, UAE-resident hosting (pilot)"],
        ["Bias", "Model less accurate for some bodies or groups", "Fairness table by acclimatization, age, BMI, language; recalibrate per group in pilot"],
        ["Explainability", "Black-box alerts are ignored", "Physics model plus Kalman filter; every alert shows estimate, 30-min forecast and reason"],
        ["Misuse", "Employer uses data to penalise slower new workers", "Tier-only employer view; contract bans use for pay, hiring or dismissal; newcomer protection (they lose most work time: 5.8 h vs 7.6 h) is flagged to regulators"],
        ["Human oversight", "Automation replaces judgement", "RED always notifies a human first-aider; emergency signs (confusion, no sweating, collapse) override the model"],
        ["Failure modes", "Sensor dropout, wrong task class, alert fatigue", f"Outlier gating, rising uncertainty widens alerts, fall back to site WBGT rule. Known weakness: {EV['warning']['amber_or_red_rate_among_never_38']:.0f}% of workers who stayed under 38 °C still saw an amber or red tier; threshold tuning and pilot feedback needed"]]
table(s, 0.3, 1.5, 9.4, rows, [1.1, 2.4, 5.9], size=8, rowh=0.44)

# ---------------------------------------------------------------- slide 10 hardware
s = S[9]
tb(s, 0.35, 1.5, 4.6, 3.9, [
    H("No custom hardware is required"),
    B([("Wearable: ", True), ("any consumer wrist or chest heart-rate band with a standard Bluetooth heart-rate profile (for example a PPG band at USD 30 to 100).", False)]),
    B([("Phone: ", True), ("a basic smartphone or feature phone with SMS or WhatsApp. Workers without a smartphone can receive voice or SMS alerts.", False)]),
    B([("Site weather: ", True), ("a public weather feed, or one low-cost WBGT sensor per site (optional, improves accuracy).", False)]),
    B([("Gateway: ", True), ("the worker's phone or a site tablet relays readings; cloud service runs the engine.", False)]),
], size=9.5)
tb(s, 5.2, 1.5, 4.5, 3.9, [
    H("Sourcing and safety"),
    B("Off-the-shelf, certified consumer devices; no implant, no ingestible and no invasive sensor in the product."),
    B("Skin-contact and battery safety follow the device vendor's certification; band choice must tolerate heat, sweat and dust."),
    B("Ingestible core-temperature pills are used only in the ethics-approved calibration sub-study, never as part of the deployed product."),
    B("Bands are issued per worker and remain the worker's property for the shift; cleaning and hygiene protocol is part of the pilot."),
    H("Classification"),
    B("This submission is software-based. Hardware is third-party and interchangeable."),
], size=9.5)

# ---------------------------------------------------------------- slide 11 pilot
s = S[10]
steps = [("Nov to Dec 2026", "Prototype, validation and GITEX presentation (this challenge)"),
         ("Q1 2027", "Recruit occupational-health physician; ethics approval; heat-chamber calibration against ingestible pill, about 30 volunteers"),
         ("Q2 2027", "Fix thresholds and alert wording; native-speaker review; privacy impact assessment; data-protection agreement"),
         ("Jun to Sep 2027", "Pilot at 1 or 2 construction or delivery sites, 50 to 100 workers, during the midday-break season; compare with matched baseline"),
         ("2028", "Scale to more employers, integrate with safety systems; evidence package for regulators and insurers")]
for k, (d, tx) in enumerate(steps):
    x = 0.3 + k * 1.9
    rect(s, x, 1.5, 1.8, 0.32, d, fill=BROWN, line=None, size=8.5, bold=True, color=RGBColor(255, 255, 255))
    rect(s, x, 1.86, 1.8, 1.05, tx, size=8)
tb(s, 0.3, 3.05, 4.6, 2.4, [
    H("Stakeholders and integration"),
    B("Abu Dhabi Department of Health and MoHRE (policy and regulatory fit)."),
    B("Contractors and logistics fleets (deployment); insurers (incentive); wearable and WhatsApp providers."),
    B("Integration: WhatsApp Business API or SMS gateway, site safety dashboards, aggregated feed for authorities."),
], size=8.8)
tb(s, 5.1, 3.05, 4.6, 2.4, [
    H("Regulatory, business and sustainability"),
    B("Positioned as wellness and safety decision support, not diagnosis; formal classification to be confirmed with the regulator before pilot."),
    B("Model: monthly fee per worker during heat season, offset by fewer lost shifts and claims; insurer-funded option."),
    B("Low cost: commodity wearables, cloud inference, no new infrastructure. Scales to any hot climate."),
], size=8.8)

# ---------------------------------------------------------------- slide 12 IP
s = S[11]
rows = [["Component", "Owner", "Licence / restriction"],
        ["HeatGuard code: simulator, estimator, policies, API, dashboard, deck, video", "Applicant (Abdul Mannaan Sayyad); written during this challenge", "No pre-existing product. Employer (Triad AI) permission and IP position being confirmed by the applicant"],
        ["Python libraries: NumPy, SciPy, Matplotlib, FastAPI, Pydantic, Uvicorn, pytest, Playwright", "Respective authors", "BSD / MIT / Apache-2.0, unrestricted use"],
        ["Fonts: Outfit, Noto Sans family, JetBrains Mono", "Respective authors", "SIL Open Font License"],
        ["Kokoro text-to-speech (video narration)", "Kokoro authors", "Apache-2.0 model and library"],
        ["Scientific equations (Gagge-type two-node balance, Stull wet-bulb, Tetens vapour pressure, DuBois body surface area)", "Published literature", "Free to implement; cited on the evidence slide"],
        ["Challenge template, Future Health and co-creator logos", "Organisers", "Used only as supplied for this submission"]]
table(s, 0.3, 1.5, 9.4, rows, [3.9, 2.4, 3.1], size=8, rowh=0.44)

# ---------------------------------------------------------------- slide 13 COI
s = S[12]
tb(s, 0.35, 1.55, 9.3, 3.5, [
    {"t": [("None declared.", True, BROWN)], "size": 14, "sp": 0},
    {"t": "The applicant has no known financial, employment, family or advisory relationship with the organisers of the Future Health Digital Innovation Challenge, its co-creators (Microsoft, Inception, MBZUAI), the collaborator, or any member of the judging panel.", "size": 11, "sp": 8},
    {"t": "The applicant works as an AI engineer at Triad AI. To the applicant's knowledge, Triad AI is not an organiser, co-creator, collaborator or sponsor of this challenge.", "size": 11, "sp": 8},
    {"t": "If a relationship becomes known before judging, the applicant will disclose it to the organisers immediately.", "size": 11, "sp": 8}])

# ---------------------------------------------------------------- slide 14 additional
s = S[13]
for sh in list(s.shapes):
    if sh.has_text_frame and sh.text_frame.text.strip() == "Additional Slide": sh._element.getparent().remove(sh._element)
tb(s, 0.3, 0.8, 9.4, 0.3, [{"t": [("Additional slide: how HeatGuard behaves on one simulated hot August shift, and what is still open", True, BROWN)], "size": 11}])
s.shapes.add_picture(f"{ROOT}/results/fig_tradeoff.png", Inches(0.3), Inches(1.2), width=Inches(4.55))
s.shapes.add_picture(f"{ROOT}/results/fig_accuracy.png", Inches(0.3), Inches(3.35), width=Inches(4.55))
demo = json.load(open(f"{ROOT}/results/demo_shift.json"))["summary"]
tb(s, 5.05, 1.2, 4.65, 4.2, [
    H("Demo shift (8 workers, 36 °C mean day, 06:00 to 18:00)"),
    B(f"Without HeatGuard, {sum(p >= 38.5 for p in demo['A_peak'])} of 8 workers reach 38.5 °C (highest {max(demo['A_peak']):.1f} °C). With HeatGuard none do (highest {max(demo['C_peak']):.1f} °C)."),
    B(f"Average productive time per worker on this very hot day: {demo['C_work_h']:.1f} h versus {demo['A_work_h']:.1f} h unprotected. The trade-off curve on the left shows how the operating point was chosen on a separate tuning set (seed 101) and then frozen."),
    B("Alerts reach each worker in Hindi, Urdu, Bengali, Malayalam, Nepali, Tagalog or English. The supervisor view hides all temperatures."),
    H("What is still open"),
    B("No human data yet. The first pilot task is calibration against ingestible-pill readings."),
    B(f"Alert burden is high in simulation ({EV['warning']['amber_or_red_rate_among_never_38']:.0f}% of never-above-38 °C workers saw at least an amber tier; red {EV['warning']['red_alert_rate_among_never_38']:.0f}%). Amber is a soft prompt; thresholds need pilot tuning."),
    B("Lowest warning rate is for acclimatized workers (76% warned 10+ min ahead)."),
    H("Try it"),
    B("Live replay console: https://abdulmannaan502.github.io/heatguard/  |  Code: https://github.com/abdulmannaan502/heatguard"),
], size=8.5)

prs.save(OUT)
print("saved", OUT)
